# -*- coding: utf-8 -*-
"""
根据实际代码生成改进后的 SDS 文档
石材幕墙污渍检测系统 软件设计规格说明文档
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

doc = Document()

# ============ 样式设置 ============
style = doc.styles['Normal']
font = style.font
font.name = '宋体'
font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

for level in range(1, 4):
    heading_style = doc.styles[f'Heading {level}']
    heading_style.font.name = '黑体'
    heading_style.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    heading_style.font.color.rgb = RGBColor(0, 0, 0)
    if level == 1:
        heading_style.font.size = Pt(18)
    elif level == 2:
        heading_style.font.size = Pt(15)
    else:
        heading_style.font.size = Pt(13)

def add_heading(text, level=1):
    doc.add_heading(text, level=level)

def add_para(text, bold=False, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = '宋体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(11)
    return p

def add_list(text, level=0):
    p = doc.add_paragraph(style='List Paragraph')
    p.paragraph_format.left_indent = Cm(1.27 + level * 0.63)
    run = p.add_run(text)
    run.font.name = '宋体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(11)
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.name = '宋体'
        run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        run.font.size = Pt(10)
    # Data rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.name = '宋体'
            run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            run.font.size = Pt(10)
    return table

# ============ 封面 ============
for _ in range(4):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('石材幕墙污渍检测系统')
run.bold = True
run.font.size = Pt(26)
run.font.name = '黑体'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('软件设计规格说明文档')
run.font.size = Pt(18)
run.font.name = '黑体'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Software Design Specification')
run.font.size = Pt(14)
run.font.italic = True

for _ in range(3):
    doc.add_paragraph()

info_lines = [
    '项目经理：于广淳（2353740）',
    '项目助理：周慧星（2351289）',
]
for line in info_lines:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(line)
    run.font.size = Pt(12)
    run.font.name = '宋体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('同济大学软件工程专业开发团队')
run.font.size = Pt(12)
run.font.name = '宋体'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('2026年6月')
run.font.size = Pt(12)
run.font.name = '宋体'
run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

doc.add_page_break()

# ============ 目录 ============
add_heading('目  录', level=1)
toc_items = [
    '1. 引言',
    '    1.1 文档目的',
    '    1.2 系统范围',
    '2. 系统概述',
    '    2.1 项目背景',
    '    2.2 系统功能概述',
    '3. 系统架构设计',
    '    3.1 架构总体设计',
    '    3.2 技术选型',
    '    3.3 前端架构',
    '    3.4 后端架构',
    '    3.5 系统集成架构',
    '4. 接口设计',
    '    4.1 API接口清单',
    '    4.2 创建检测任务 POST /api/detections',
    '    4.3 查询检测列表 GET /api/detections',
    '    4.4 获取检测详情 GET /api/detections/{task_id}',
    '    4.5 重试检测 POST /api/detections/{task_id}/retry',
    '    4.6 获取签名URL GET /api/detections/{task_id}/signed-url',
    '    4.7 删除检测 DELETE /api/detections/{task_id}',
    '    4.8 批量删除 DELETE /api/detections/batch/delete',
    '5. 数据设计',
    '    5.1 数据库设计',
    '    5.2 数据字典',
    '    5.3 数据流图',
    '6. 组件设计',
    '    6.1 后端组件',
    '    6.2 前端组件',
    '    6.3 组件依赖关系',
    '7. 用户界面设计',
    '    7.1 系统整体布局',
    '    7.2 检测控制台页面',
    '    7.3 历史记录页面',
    '    7.4 检测报告导出',
    '8. 安全设计',
    '    8.1 身份认证',
    '    8.2 权限控制',
    '9. 部署设计',
    '    9.1 后端部署',
    '    9.2 前端部署',
    '    9.3 环境变量配置',
    '10. 附录',
    '    10.1 缩略语与术语表',
    '    10.2 参考文档',
]
for item in toc_items:
    p = doc.add_paragraph()
    run = p.add_run(item)
    run.font.name = '宋体'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    run.font.size = Pt(11)

doc.add_page_break()

# ============ 1. 引言 ============
add_heading('1. 引言', level=1)

add_heading('1.1 文档目的', level=2)
add_para('本文档旨在为石材幕墙污渍检测系统提供完整的软件设计规格说明，涵盖系统架构设计、组件设计、数据设计、接口设计、用户界面设计、安全设计及部署设计，作为开发团队实施和维护的技术依据，同时也为项目评审与验收提供参考。', indent=True)
add_para('本系统基于YOLO系列深度学习模型与图像处理算法，采用前后端分离架构，实现石材幕墙污渍的自动检测、面积计算、结果可视化及历史记录管理等功能。系统前端集成于已有的智慧幕墙管理平台（基于Nuxt.js 3），后端独立部署为FastAPI微服务。', indent=True)

add_heading('1.2 系统范围', level=2)
add_para('石材幕墙污渍检测系统的主要客户及终端用户为同济大学土木工程团队管理层、物业管理公司及建筑维护公司。')
add_para('本系统作为同济大学计算机科学与技术学院智慧幕墙项目的子系统，专注于实现以下功能范围：')
add_list('石材幕墙污渍检测模块的开发与集成（基于YOLO系列模型，支持本地推理与云端API推理双模式）')
add_list('污渍区域检测与分类（轻度light、中度moderate、重度severe三级分类）')
add_list('污渍占比计算与整体清洁度评估')
add_list('前端检测控制台与历史记录页面的开发，集成于已有智慧幕墙管理平台')
add_list('基于MySQL和OSS的数据存储与管理（4张数据表）')
add_list('通过Docker容器化技术实现后端服务部署')
add_para('本系统不涉及硬件设备的采购与维护。前端作为子模块集成于已有智慧幕墙项目前端，训练所得YOLO系列模型可供复用。', indent=True)

doc.add_page_break()

# ============ 2. 系统概述 ============
add_heading('2. 系统概述', level=1)

add_heading('2.1 项目背景', level=2)
add_para('随着现代建筑中石材幕墙的广泛应用，其表面的污渍、粉尘和污染物问题日益突出。石材幕墙不仅具有美观的视觉效果，还发挥着隔热、隔音、抗风压等重要功能。然而，受气候变化、环境污染及建筑使用等因素影响，幕墙表面极易积累污渍，不仅影响建筑美观，还可能损害其使用寿命和功能。', indent=True)
add_para('目前，传统的幕墙检测主要依赖人工巡检，存在效率低下、成本高昂、高空作业安全风险大等明显缺陷。基于此，本项目利用计算机视觉与深度学习技术，开发一套智能石材幕墙污渍检测系统，实现对石材幕墙污渍的快速、精准检测，以达到以下目标：')
add_list('提升检测效率：将人工巡检转变为自动化检测，大幅缩短检测时间')
add_list('降低检测成本：减少对专业巡检人员的依赖，降低人力成本')
add_list('提高安全性：避免人工高空作业，降低安全风险')
add_list('提供数据支持：为建筑的长期维护和清洁计划提供科学、量化的数据依据')

add_heading('2.2 系统功能概述', level=2)

add_heading('2.2.1 图像上传与管理功能', level=3)
add_para('（1）图像上传：用户可通过系统上传建筑幕墙图片，支持拖拽上传、文件选择上传及批量上传三种方式，支持JPEG、PNG等主流图像格式，单文件大小不超过50MB。系统会对超过100KB的图像自动进行压缩处理以加快上传速度。')
add_para('（2）图像命名与元数据录入：上传过程中弹出命名对话框，用户需为图像指定名称，并填写建筑名称、楼层、分区、描述等元数据信息，便于后续管理与检索。')
add_para('（3）待检测队列管理：成功上传的图片自动进入待检测队列列表，支持单独上传、批量上传、删除、切换推理模式（本地/云端）等操作。')
add_para('（4）检测图片选择：用户可从待检测队列中选取指定图片进行检测，也可一键批量提交所有待检测图片。')

add_heading('2.2.2 幕墙污渍检测功能', level=3)
add_para('（1）污渍检测与分类：系统调用训练好的YOLO系列模型对图像进行目标检测，直接识别图像中的污渍区域，并按严重程度分为light（轻度）、moderate（中度）、severe（重度）三级。')
add_para('（2）双模式推理：系统支持本地模型推理和云端API推理两种模式。本地模式使用Ultralytics YOLO加载本地best.pt模型文件；云端模式通过HTTP请求调用远程推理服务。')
add_para('（3）智能同步/异步处理：图像文件大小不超过2MB时采用同步处理（直接返回结果），超过2MB时采用异步后台处理（通过FastAPI BackgroundTasks）。')
add_para('（4）结果可视化：系统将检测结果（带标注的图像、污渍类型、污渍面积占比、整体清洁度等）返回至前端界面进行可视化展示，包括原图与检测后图片的对比展示。')

add_heading('2.2.3 检测报告导出功能', level=3)
add_para('（1）当次报告导出：在检测控制台页面完成单次检测后，用户可点击【导出PDF】按钮，以PDF格式导出本次检测报告，内容包括任务信息、建筑信息、污渍类型、污渍占比、区域明细、原图与检测后图片等。')
add_para('（2）历史报告导出：在历史记录查询页面，用户可针对任意一条历史检测记录点击【导出PDF】按钮，导出该条记录的检测报告。')
add_para('（3）批量报告导出：用户可勾选多条历史记录，点击【导出PDF】按钮批量导出多条检测记录的报告。')

add_heading('2.2.4 历史记录查询功能', level=3)
add_para('（1）记录存储与归档：系统将每次污渍检测结果自动保存为历史记录，包含检测时间、建筑名称、楼层、分区、图片名称、污渍类型、污渍占比、整体清洁度等信息。')
add_para('（2）多维筛选：用户可按建筑名称、检测状态（待处理/处理中/已完成/失败）、时间范围对历史记录进行过滤。')
add_para('（3）分页展示：历史记录按检测时间倒序排列，分页显示，每页默认10条记录，支持翻页浏览。')
add_para('（4）详情查看：点击历史记录可打开抽屉式详情面板，查看完整检测信息，包括原图与检测后图片的对比、区域明细表格（污渍类型、置信度、边界框坐标）等。')
add_para('（5）重试与删除：对失败的检测任务支持一键重试；支持单条删除和批量删除历史记录。')
add_para('（6）客户端缓存：历史列表采用15分钟客户端缓存机制，减少重复请求，提升响应速度。')

doc.add_page_break()

# ============ 3. 系统架构设计 ============
add_heading('3. 系统架构设计', level=1)

add_heading('3.1 架构总体设计', level=2)
add_para('本系统采用前后端分离的微服务架构。污渍检测功能作为智慧幕墙管理平台的一个子模块，前端页面集成于已有平台前端，后端独立部署为FastAPI微服务。整体架构如下：')
add_list('前端层：基于Nuxt.js 3框架（Vue 3 + SSR可选），集成于智慧幕墙管理平台，负责图像上传、检测发起、结果展示与报告下载等用户操作，通过RESTful API与后端通信。')
add_list('代理层：Nuxt.js Nitro服务器中间件作为反向代理，将前端请求路由到多个后端服务（污渍检测、用户认证、玻璃检测、腐蚀检测等）。')
add_list('后端服务层：基于FastAPI框架提供污渍检测业务逻辑处理和API服务，包括图像接收、YOLO模型推理、结果处理及数据持久化等核心功能。')
add_list('算法层：集成YOLO系列深度学习模型（Ultralytics YOLO），支持本地推理（best.pt）和云端API推理两种模式，实现污渍检测与三级分类。')
add_list('数据存储层：采用MySQL数据库（4张数据表）存储检测任务元数据与结果，采用OSS对象存储（MinIO兼容）保存原始图像及检测结果图像。')
add_list('部署层：后端通过Docker容器化技术部署，前端通过Nuxt.js内置服务器或静态生成方式部署。')

add_heading('3.2 技术选型', level=2)
add_para('各层技术选型如下表所示：')
add_table(
    ['层次', '技术', '版本/说明'],
    [
        ['前端框架', 'Nuxt.js 3', '3.17.3，基于Vue 3.5.13，CSR模式（ssr: false）'],
        ['UI组件库', 'Element Plus + Nuxt UI Pro', '提供表格、表单、对话框、抽屉等组件'],
        ['CSS框架', 'Tailwind CSS', '通过Nuxt UI集成'],
        ['HTTP客户端', 'Axios', '前后端API通信'],
        ['图表库', 'ECharts 5.6.0', '数据可视化'],
        ['PDF生成', 'jsPDF + html2canvas', '客户端PDF报告导出'],
        ['状态管理', 'Pinia + useState', '前端状态管理'],
        ['后端框架', 'FastAPI', '0.116.1，Python异步Web框架'],
        ['ASGI服务器', 'Uvicorn', '0.35.0'],
        ['数据库', 'MySQL 8.0+', '通过PyMySQL直连，无ORM'],
        ['对象存储', 'OSS（MinIO兼容）', '图像文件存储'],
        ['ML模型', 'Ultralytics YOLO', '本地best.pt模型推理'],
        ['图像处理', 'Pillow + NumPy', '图像读取、绘制标注'],
        ['认证', 'PyJWT', 'HS256 Bearer Token'],
        ['数据验证', 'Pydantic', '2.11.7，请求/响应数据校验'],
        ['容器化', 'Docker + Docker Compose', '后端服务部署'],
    ]
)

add_heading('3.3 前端架构', level=2)
add_para('前端采用Nuxt.js 3框架，以CSR（客户端渲染）模式运行。污渍检测模块的前端页面作为子模块集成于已有的智慧幕墙管理平台前端。')

add_heading('3.3.1 前端目录结构', level=3)
add_para('污渍检测相关的核心文件：')
add_table(
    ['路径', '说明'],
    [
        ['pages/stonedirty/detection.vue', '检测控制台页面（新版，含队列管理）'],
        ['pages/stonedirty/history.vue', '历史记录查询页面'],
        ['pages/stonedirty/mainpage.vue', '旧版检测页面（兼容保留）'],
        ['api/detections.ts', '新后端API封装（调用FastAPI服务）'],
        ['api/stain.js', '旧版API封装（调用遗留服务）'],
        ['types/detection.ts', 'TypeScript类型定义'],
        ['composables/usePendingQueue.ts', '待检测队列状态管理'],
        ['utils/detectionPdfExport.ts', 'PDF报告导出工具'],
        ['utils/compressImage.ts', '图像压缩工具'],
        ['layouts/default.vue', '主布局（含侧边栏导航）'],
        ['middleware/auth.global.js', '全局认证中间件'],
        ['server/middleware/proxy.ts', 'Nitro反向代理中间件'],
    ]
)

add_heading('3.3.2 前端代理层', level=3)
add_para('Nuxt.js Nitro服务器中间件负责将前端请求路由到对应的后端服务：')
add_table(
    ['前端路径前缀', '目标后端', '用途'],
    [
        ['/detection-api/*', 'localhost:8081', '新版FastAPI污渍检测后端'],
        ['/api/auth/*', 'Nitro本地API', '登录、注册、登出、用户信息'],
        ['/api/*', '8.159.143.133:8000', '用户账户/认证服务'],
        ['/api/detect/*', '8.153.161.229:8003', '玻璃裂纹/平整度检测'],
        ['/api/corrosion/*', 'Nitro本地API', '腐蚀检测'],
        ['/predict, /history', '47.102.208.89:8007', '遗留污渍检测服务'],
        ['/oss/*', '8.159.143.133:9000', 'OSS对象存储'],
    ]
)

add_heading('3.4 后端架构', level=2)
add_para('后端采用FastAPI框架，独立部署于8081端口，提供污渍检测的完整RESTful API。')

add_heading('3.4.1 后端目录结构', level=3)
add_table(
    ['路径', '说明'],
    [
        ['app/main.py', 'FastAPI应用入口，CORS配置，路由注册'],
        ['app/api/detections.py', '检测API路由（CRUD + 重试 + 签名URL）'],
        ['app/api/health.py', '健康检查端点'],
        ['app/core/config.py', '环境变量配置（Settings数据类）'],
        ['app/core/auth.py', 'JWT认证（PyJWT HS256解码）'],
        ['app/core/response.py', '统一JSON响应格式（ok/fail）'],
        ['app/schemas/detection.py', 'Pydantic数据模型'],
        ['app/services/model_adapter.py', 'YOLO模型适配器（本地+云端推理）'],
        ['app/services/mysql_repo.py', 'MySQL数据仓库（PyMySQL直连）'],
        ['app/workers/tasks.py', '后台异步检测任务'],
        ['sql/final_detection_schema.sql', 'MySQL数据库Schema'],
    ]
)

add_heading('3.4.2 后端处理流程', level=3)
add_para('（1）同步处理流程（图像 ≤ 2MB）：')
add_list('1. 接收multipart/form-data请求（image文件 + 元数据表单字段）')
add_list('2. 验证JWT Token，提取user_id')
add_list('3. 验证图像格式（仅支持image/jpeg、image/png）')
add_list('4. 将原始图像上传至OSS，获取存储路径')
add_list('5. 在MySQL中创建inspection_tasks、task_images、detection_results三条记录（状态为pending）')
add_list('6. 更新任务状态为processing')
add_list('7. 调用model_adapter执行YOLO推理（本地或云端）')
add_list('8. 将检测后图片上传至OSS，更新task_images.processed_image_path')
add_list('9. 更新inspection_tasks（污渍检测结果）、detection_results（指标）、result_regions（区域明细）')
add_list('10. 返回完整的任务详情JSON')

add_para('（2）异步处理流程（图像 > 2MB）：')
add_list('1. 步骤1-5与同步流程相同')
add_list('2. 通过FastAPI BackgroundTasks将检测任务提交到后台线程')
add_list('3. 立即返回任务详情（状态为pending）')
add_list('4. 后台线程执行步骤6-9')
add_list('5. 前端通过轮询（每2秒）查询任务状态直到完成或失败')

add_heading('3.5 系统集成架构', level=2)
add_para('污渍检测系统作为智慧幕墙管理平台的子模块，与其他模块共享前端框架和用户认证体系：')
add_table(
    ['模块', '前端路由', '后端服务', '说明'],
    [
        ['污渍检测', '/stonedirty/*', 'localhost:8081 (FastAPI)', '本文档描述的系统'],
        ['玻璃裂纹检测', '/glass-inspection/crack', '8.153.161.229:8003', '独立后端服务'],
        ['玻璃平整度检测', '/glass-inspection/flatness', '8.153.161.229:8003', '独立后端服务'],
        ['腐蚀检测', '/corrosion/*', 'Nitro本地API', 'Nuxt服务器端处理'],
        ['振动监测', '/vibration/*', '远程服务', '独立后端服务'],
        ['性能评估', '/resilience/*', '远程服务', '独立后端服务'],
        ['用户管理', '/userManage, /accountManagement', '8.159.143.133:8000', '统一认证服务'],
    ]
)

doc.add_page_break()

# ============ 4. 接口设计 ============
add_heading('4. 接口设计', level=1)
add_para('本系统后端基于FastAPI框架，对外暴露RESTful API接口，所有接口均需携带JWT Bearer Token进行身份认证。统一响应格式为：')
add_para('{"code": 0, "data": ..., "message": "..."}（code=0表示成功，非0表示失败）')

add_heading('4.1 API接口清单', level=2)
add_table(
    ['方法', '路径', '说明', '认证'],
    [
        ['GET', '/api/health', '健康检查', '否'],
        ['POST', '/api/detections', '创建检测任务', '是'],
        ['GET', '/api/detections', '查询检测列表', '是'],
        ['GET', '/api/detections/{task_id}', '获取检测详情', '是'],
        ['POST', '/api/detections/{task_id}/retry', '重试失败任务', '是'],
        ['GET', '/api/detections/{task_id}/signed-url', '获取图像签名URL', '是'],
        ['DELETE', '/api/detections/{task_id}', '删除单条检测', '是'],
        ['DELETE', '/api/detections/batch/delete', '批量删除检测', '是'],
    ]
)

add_heading('4.2 创建检测任务 POST /api/detections', level=2)
add_para('接口说明：接收用户上传的图像文件及元数据，创建检测任务并执行污渍检测。小文件（≤2MB）同步返回结果，大文件异步处理。')
add_para('请求方式：POST，Content-Type: multipart/form-data')
add_para('请求参数：')
add_table(
    ['参数名', '类型', '必填', '说明'],
    [
        ['image', 'File', '是', '图像文件（仅支持image/jpeg、image/png）'],
        ['building_name', 'string', '是', '建筑名称'],
        ['location_floor', 'int', '否', '楼层号'],
        ['location_section', 'string', '否', '分区名称'],
        ['description', 'string', '否', '描述信息'],
        ['inference_mode', 'string', '否', '推理模式："local"（默认）或"cloud"'],
    ]
)
add_para('返回参数（data对象）：')
add_table(
    ['字段', '类型', '说明'],
    [
        ['id', 'string', '任务ID（自增主键）'],
        ['userId', 'string', '用户ID'],
        ['buildingName', 'string', '建筑名称'],
        ['locationFloor', 'int|null', '楼层'],
        ['locationSection', 'string|null', '分区'],
        ['status', 'string', '任务状态：pending/processing/done/failed'],
        ['createdAt', 'datetime', '创建时间'],
        ['updatedAt', 'datetime', '更新时间'],
        ['imageName', 'string', '图片名称'],
        ['imagePath', 'string', 'OSS原始图像路径'],
        ['imageSignedUrl', 'string', '原始图像签名访问URL'],
        ['processedImagePath', 'string|null', 'OSS检测后图像路径'],
        ['processedImageSignedUrl', 'string|null', '检测后图像签名URL'],
        ['summary', 'string|null', '检测结果摘要'],
        ['stainDetected', 'boolean|null', '是否检测到污渍'],
        ['stainType', 'string|null', '主要污渍类型：light/moderate/severe'],
        ['affectedAreaPercentage', 'number|null', '污渍面积占比（%）'],
        ['regions', 'array', '检测区域明细列表'],
        ['metrics', 'object|null', '检测指标（runtimeMs, overallCleanliness, inferenceMode）'],
        ['errorMessage', 'string|null', '错误信息（失败时）'],
    ]
)

add_heading('4.3 查询检测列表 GET /api/detections', level=2)
add_para('接口说明：分页查询当前用户的检测任务列表，支持多维筛选。')
add_para('请求参数（Query）：')
add_table(
    ['参数名', '类型', '必填', '说明'],
    [
        ['current_page', 'int', '否', '页码，默认1，最小1'],
        ['size', 'int', '否', '每页条数，默认10，范围1-100'],
        ['status', 'string', '否', '状态筛选：pending/processing/done/failed'],
        ['building_name', 'string', '否', '建筑名称模糊搜索'],
        ['start_time', 'string', '否', '开始时间（ISO格式）'],
        ['end_time', 'string', '否', '结束时间（ISO格式）'],
    ]
)
add_para('返回参数：')
add_table(
    ['字段', '类型', '说明'],
    [
        ['items', 'array', '检测任务列表（字段同4.2，但regions为空数组，imageSignedUrl等为null）'],
        ['pagination.current_page', 'int', '当前页码'],
        ['pagination.size', 'int', '每页条数'],
        ['pagination.total', 'int', '总记录数'],
    ]
)

add_heading('4.4 获取检测详情 GET /api/detections/{task_id}', level=2)
add_para('接口说明：获取指定检测任务的完整详情，包括区域明细和签名URL。')
add_para('路径参数：task_id（任务ID）')
add_para('返回参数：完整的任务详情对象（字段同4.2）。')

add_heading('4.5 重试检测 POST /api/detections/{task_id}/retry', level=2)
add_para('接口说明：对失败的检测任务重新提交执行。将任务状态重置为pending，从OSS重新下载图像并提交后台检测。')
add_para('路径参数：task_id（任务ID）')
add_para('返回参数：{"id": task_id, "status": "pending"}')

add_heading('4.6 获取签名URL GET /api/detections/{task_id}/signed-url', level=2)
add_para('接口说明：获取检测任务的原始图像和检测后图像的签名访问URL。')
add_para('返回参数：')
add_table(
    ['字段', '类型', '说明'],
    [
        ['id', 'string', '任务ID'],
        ['imagePath', 'string', 'OSS存储路径'],
        ['imageSignedUrl', 'string', '签名访问URL'],
    ]
)

add_heading('4.7 删除检测 DELETE /api/detections/{task_id}', level=2)
add_para('接口说明：删除指定检测任务及其关联的所有数据（任务、图像记录、检测结果、区域明细）。')
add_para('路径参数：task_id（任务ID）')

add_heading('4.8 批量删除 DELETE /api/detections/batch/delete', level=2)
add_para('接口说明：批量删除多条检测记录。')
add_para('请求参数：task_ids（Query参数，逗号分隔的任务ID列表）')

doc.add_page_break()

# ============ 5. 数据设计 ============
add_heading('5. 数据设计', level=1)

add_heading('5.1 数据库设计', level=2)
add_para('本系统采用双存储架构：')
add_list('MySQL关系型数据库（库名：stain，字符集utf8mb4）：存储检测任务元数据、图像记录、检测结果和区域明细，共4张数据表。')
add_list('OSS对象存储（MinIO兼容）：存储用户上传的原始图像文件及检测结果图像文件，通过签名URL进行访问。')

add_heading('5.2 数据字典', level=2)

add_heading('5.2.1 检测任务表（inspection_tasks）', level=3)
add_para('表描述：存储每次检测任务的主信息。')
add_table(
    ['字段', '类型', '约束', '说明'],
    [
        ['id', 'bigint unsigned', 'PK, AUTO_INCREMENT', '任务ID'],
        ['user_id', 'varchar(64)', 'NOT NULL', '用户ID（来自JWT）'],
        ['building_name', 'varchar(255)', 'NOT NULL', '建筑名称'],
        ['location_floor', 'int', 'NULL', '楼层号'],
        ['location_section', 'varchar(255)', 'NULL', '分区名称'],
        ['description', 'text', 'NULL', '描述信息'],
        ['status', 'enum', 'NOT NULL, DEFAULT pending', '状态：pending/processing/done/failed'],
        ['summary', 'text', 'NULL', '检测结果摘要'],
        ['stain_detected', 'tinyint(1)', 'NULL', '是否检测到污渍（0/1）'],
        ['stain_type', 'varchar(255)', 'NULL', '主要污渍类型'],
        ['affected_area_percentage', 'decimal(6,2)', 'NULL', '污渍面积占比（%）'],
        ['error_message', 'text', 'NULL', '错误信息'],
        ['created_at', 'datetime', 'NOT NULL', '创建时间'],
        ['updated_at', 'datetime', 'NOT NULL', '更新时间'],
    ]
)
add_para('索引：')
add_list('idx_inspection_tasks_user_created (user_id, created_at DESC)')
add_list('idx_inspection_tasks_user_status_created (user_id, status, created_at DESC)')
add_list('idx_inspection_tasks_status (status)')
add_list('idx_inspection_tasks_building_name (building_name)')

add_heading('5.2.2 任务图像表（task_images）', level=3)
add_para('表描述：存储检测任务关联的图像信息。')
add_table(
    ['字段', '类型', '约束', '说明'],
    [
        ['id', 'bigint unsigned', 'PK, AUTO_INCREMENT', '记录ID'],
        ['task_id', 'bigint unsigned', 'FK -> inspection_tasks.id (CASCADE)', '关联任务ID'],
        ['user_id', 'varchar(64)', 'NOT NULL', '用户ID'],
        ['image_name', 'varchar(255)', 'NULL', '图片名称'],
        ['original_image_path', 'varchar(500)', 'NOT NULL', 'OSS原始图像路径'],
        ['processed_image_path', 'varchar(500)', 'NULL', 'OSS检测后图像路径'],
        ['mime_type', 'varchar(100)', 'NULL', 'MIME类型'],
        ['file_size', 'bigint unsigned', 'NULL', '文件大小（字节）'],
        ['created_at', 'datetime', 'NOT NULL', '创建时间'],
    ]
)

add_heading('5.2.3 检测结果表（detection_results）', level=3)
add_para('表描述：存储检测任务的结果元数据和运行指标。')
add_table(
    ['字段', '类型', '约束', '说明'],
    [
        ['id', 'bigint unsigned', 'PK, AUTO_INCREMENT', '记录ID'],
        ['task_id', 'bigint unsigned', 'FK, UNIQUE (CASCADE)', '关联任务ID（一对一）'],
        ['user_id', 'varchar(64)', 'NOT NULL', '用户ID'],
        ['status', 'enum', 'NOT NULL, DEFAULT pending', '状态'],
        ['summary', 'text', 'NULL', '检测摘要'],
        ['metrics', 'json', 'NULL', 'JSON指标（runtimeMs, overallCleanliness, inferenceMode）'],
        ['error_message', 'text', 'NULL', '错误信息'],
        ['created_at', 'datetime', 'NOT NULL', '创建时间'],
        ['updated_at', 'datetime', 'NOT NULL', '更新时间'],
        ['processed_at', 'datetime', 'NULL', '处理完成时间'],
    ]
)

add_heading('5.2.4 检测区域表（result_regions）', level=3)
add_para('表描述：存储检测到的每个污渍区域的详细信息。')
add_table(
    ['字段', '类型', '约束', '说明'],
    [
        ['id', 'bigint unsigned', 'PK, AUTO_INCREMENT', '记录ID'],
        ['task_id', 'bigint unsigned', 'FK -> inspection_tasks.id (CASCADE)', '关联任务ID'],
        ['label', 'varchar(255)', 'NOT NULL', '污渍类型：light/moderate/severe'],
        ['confidence', 'decimal(5,4)', 'NOT NULL', '置信度（0-1）'],
        ['x1', 'decimal(8,5)', 'NOT NULL', '边界框左上角X（归一化坐标）'],
        ['y1', 'decimal(8,5)', 'NOT NULL', '边界框左上角Y（归一化坐标）'],
        ['x2', 'decimal(8,5)', 'NOT NULL', '边界框右下角X（归一化坐标）'],
        ['y2', 'decimal(8,5)', 'NOT NULL', '边界框右下角Y（归一化坐标）'],
        ['created_at', 'datetime', 'NOT NULL', '创建时间'],
    ]
)

add_heading('5.3 数据流图', level=2)

add_heading('5.3.1 污渍检测功能数据流图', level=3)
add_list('1. 用户通过前端选择图片，填写建筑名称等元数据，提交至后端（POST /api/detections，multipart/form-data）。')
add_list('2. 后端接收图像文件，上传至OSS对象存储，获取原始图像存储路径。')
add_list('3. 后端在MySQL中创建inspection_tasks、task_images、detection_results记录。')
add_list('4. 后端调用model_adapter执行YOLO推理（本地best.pt或云端API），获取检测结果（污渍区域、类型、置信度、边界框）。')
add_list('5. 后端将检测后标注图像上传至OSS，更新task_images.processed_image_path。')
add_list('6. 后端更新MySQL中的检测结果（inspection_tasks汇总、detection_results指标、result_regions区域明细）。')
add_list('7. 后端将结构化JSON结果返回前端。')
add_list('8. 前端解析JSON数据，展示检测结果（原图对比、污渍类型、占比、区域明细表格）。')

add_heading('5.3.2 历史记录查询数据流图', level=3)
add_list('1. 用户在前端填写查询条件（建筑名称、状态、时间范围），提交查询请求至后端。')
add_list('2. 后端根据查询条件从MySQL数据库中分页检索历史记录（inspection_tasks关联task_images、detection_results）。')
add_list('3. 后端将查询结果（含图像路径、检测指标）以JSON格式返回前端。')
add_list('4. 前端展示历史列表；用户点击详情时，通过signed-url接口获取图像签名URL，从OSS加载并展示图像。')

doc.add_page_break()

# ============ 6. 组件设计 ============
add_heading('6. 组件设计', level=1)

add_heading('6.1 后端组件', level=2)

add_heading('6.1.1 应用入口（app/main.py）', level=3)
add_para('FastAPI应用入口，负责：创建FastAPI实例、配置CORS中间件（允许跨域请求）、注册health和detections两个API路由。')

add_heading('6.1.2 API路由层（app/api/）', level=3)
add_para('detections.py：定义检测相关的7个RESTful API端点，处理请求参数验证、文件格式校验，调用服务层执行业务逻辑。')
add_para('health.py：提供健康检查端点，返回服务状态、时间戳和版本号。')

add_heading('6.1.3 核心层（app/core/）', level=3)
add_para('config.py：基于dataclass的配置管理，从环境变量加载所有配置项（MySQL连接、OSS地址、JWT密钥、YOLO模型路径、云端模型配置等）。')
add_para('auth.py：JWT认证模块，使用PyJWT库解码HS256 Bearer Token，从payload中提取user_id（支持sub/id/email字段）。')
add_para('response.py：统一JSON响应格式封装（ok/fail辅助函数）。')

add_heading('6.1.4 服务层（app/services/）', level=3)
add_para('model_adapter.py：YOLO模型适配器，核心检测组件。')
add_list('支持本地推理（_detect_local）和云端推理（_detect_cloud）两种模式')
add_list('本地模式：加载best.pt模型，执行predict，解析boxes和masks，生成标注图像')
add_list('云端模式：通过HTTP POST调用远程推理API，解析返回的JSON结果，使用Pillow绘制标注')
add_list('污渍分类：light（轻度，浅蓝）、moderate（中度，青色）、severe（重度，深蓝）')
add_list('计算指标：affected_area_percentage（污渍占比）、overall_cleanliness（整体清洁度 = 100 - 污渍占比）')
add_para('mysql_repo.py：MySQL数据仓库，封装所有数据库操作。')
add_list('使用PyMySQL直连MySQL，每次操作独立创建连接（无连接池）')
add_list('封装CRUD操作：create_detection、update_detection_processing/done/failed、list_tasks、get_task_detail、delete_task等')
add_list('OSS图像上传：upload_image、upload_processed_image（支持重试机制）')
add_list('签名URL生成：create_signed_image_url')

add_heading('6.1.5 后台任务层（app/workers/）', level=3)
add_para('tasks.py：后台异步检测任务。')
add_list('process_detection_task：接收图像字节数据，执行同步推理（用于FastAPI BackgroundTasks）')
add_list('process_detection_task_from_storage：从OSS下载图像后执行推理（用于重试场景）')

add_heading('6.2 前端组件', level=2)

add_heading('6.2.1 检测控制台（pages/stonedirty/detection.vue）', level=3)
add_para('新版检测页面，采用左右分栏布局：')
add_list('左侧：上传图片卡片（拖拽/点击上传）+ 待检测队列面板（PendingQueuePanelImproved组件）')
add_list('右侧：检测结果展示卡片（任务信息表格、区域明细表格、原图与检测后图片对比、PDF导出）')
add_list('支持批量上传、队列管理、单个/批量提交检测、轮询状态更新')

add_heading('6.2.2 历史记录（pages/stonedirty/history.vue）', level=3)
add_para('历史记录查询页面：')
add_list('搜索栏：建筑名称输入、状态下拉筛选、时间范围选择器')
add_list('操作按钮：查询、重置、刷新、导出PDF、批量删除')
add_list('数据表格：任务ID、图片名称、建筑、状态、推理模式、总结、污渍占比、创建时间、操作列')
add_list('详情抽屉：el-drawer组件，展示完整检测信息、区域明细表格、原图与检测后图片')
add_list('客户端缓存：15分钟Map缓存，减少重复请求')

add_heading('6.2.3 API封装层（api/detections.ts）', level=3)
add_para('封装所有与FastAPI后端的通信：')
add_list('createDetectionTask：POST multipart/form-data创建检测任务')
add_list('getDetectionTask：GET获取任务详情')
add_list('getDetectionList：GET分页查询列表')
add_list('getDetectionSignedUrl：GET获取图像签名URL')
add_list('retryDetection：POST重试失败任务')
add_list('deleteDetection：DELETE删除单条记录')
add_list('batchDeleteDetections：DELETE批量删除')
add_list('所有请求自动携带localStorage中的authToken作为Bearer Token')

add_heading('6.3 组件依赖关系', level=2)
add_para('后端组件依赖关系：')
add_list('API路由层 → 服务层（model_adapter, mysql_repo）')
add_list('API路由层 → 核心层（auth, config, response）')
add_list('服务层（mysql_repo） → 核心层（config）')
add_list('服务层（model_adapter） → 核心层（config）')
add_list('后台任务层 → 服务层（model_adapter, mysql_repo）')
add_para('前端组件依赖关系：')
add_list('页面组件 → API封装层 → Axios → Nitro代理 → FastAPI后端')
add_list('页面组件 → Composables（usePendingQueue） → 响应式状态')
add_list('页面组件 → Utils（detectionPdfExport, compressImage）')
add_list('全局中间件（auth.global.js） → localStorage中的authToken')

doc.add_page_break()

# ============ 7. 用户界面设计 ============
add_heading('7. 用户界面设计', level=1)
add_para('本系统前端基于Nuxt.js 3框架开发，遵循简洁直观的设计原则，确保用户可高效完成图像上传、污渍检测、结果查看、报告下载及历史查询等核心操作流程。污渍检测模块作为子模块集成于智慧幕墙管理平台，通过侧边栏导航进入。')

add_heading('7.1 系统整体布局', level=2)
add_para('系统采用侧边栏 + 主内容区的经典后台管理布局：')
add_list('侧边栏：包含所有功能模块的导航入口（振动监测、玻璃裂纹检测、玻璃平整度检测、性能评估、污渍检测、腐蚀检测等）')
add_list('主内容区：根据当前路由渲染对应的页面组件')
add_list('全局认证：未登录用户自动跳转至登录页面；权限控制基于localStorage中的权限字段（access_system_a ~ access_system_z, is_superuser）')

add_heading('7.2 检测控制台页面', level=2)
add_para('检测控制台（/stonedirty/detection）是污渍检测的核心操作页面，采用左右分栏布局。')

add_heading('7.2.1 左侧操作区', level=3)
add_para('（1）上传图片卡片：')
add_list('支持拖拽上传和点击上传两种方式')
add_list('支持同时选择多张图片批量上传')
add_list('支持JPEG、PNG格式，单文件不超过50MB')
add_list('超过100KB的图片自动压缩后上传')
add_para('（2）待检测队列面板：')
add_list('显示所有已上传待检测的图片列表')
add_list('每项显示：图片预览缩略图、自定义名称、上传状态、进度条')
add_list('支持操作：单独上传、删除、切换推理模式（本地模型/云端模型）')
add_list('批量操作：一键上传全部、清除已完成项')
add_list('上传完成后可直接点击查看检测结果')

add_heading('7.2.2 右侧展示区', level=3)
add_para('检测结果卡片：')
add_list('任务信息表格：任务ID、图片名称、状态（彩色标签）、推理模式、建筑、楼层、分区、污渍类型、污渍占比、检测时间、总结')
add_list('区域明细表格：每个检测到的污渍区域的类型、置信度（4位小数）、边界框坐标(x1,y1,x2,y2)')
add_list('错误提示：失败时显示错误信息Alert')
add_list('图像对比展示：原图与检测后图片并排显示，支持点击预览大图')
add_list('操作按钮：刷新（重新获取签名URL）、导出PDF')

add_heading('7.3 历史记录页面', level=2)
add_para('历史记录页面（/stonedirty/history）提供检测历史的查询、浏览和管理功能。')

add_heading('7.3.1 搜索与筛选', level=3)
add_list('建筑名称：文本输入，支持模糊搜索')
add_list('状态下拉：全部/待处理/处理中/已完成/失败')
add_list('时间范围：日期时间范围选择器')
add_list('操作按钮：查询、重置、刷新、导出PDF（批量）、删除（批量）')

add_heading('7.3.2 数据表格', level=3)
add_list('支持多选（checkbox列）')
add_list('显示字段：任务ID、图片名称、建筑、状态（彩色标签）、推理模式、总结、污渍占比、创建时间')
add_list('操作列：详情、导出PDF、重试（仅失败状态）、删除')
add_list('分页：底部Pagination组件，显示总数')

add_heading('7.3.3 详情抽屉', level=3)
add_list('从右侧滑出的el-drawer组件，宽度50%')
add_list('任务信息描述列表（el-descriptions）')
add_list('区域明细表格')
add_list('错误信息Alert（失败时）')
add_list('原图与检测后图片对比展示')
add_list('定时刷新签名URL（每1小时）')

add_heading('7.4 检测报告导出', level=2)
add_para('系统支持PDF格式的检测报告导出，使用jsPDF + html2canvas在客户端生成PDF文件。')
add_list('检测控制台页面：导出当前查看的单条检测结果报告')
add_list('历史记录页面：导出单条历史记录报告，或批量导出勾选的多条记录报告')
add_list('报告内容包括：任务信息、建筑信息、污渍类型、污渍占比、区域明细、原图与检测后图片等')

doc.add_page_break()

# ============ 8. 安全设计 ============
add_heading('8. 安全设计', level=1)

add_heading('8.1 身份认证', level=2)
add_para('系统采用JWT（JSON Web Token）Bearer Token认证机制：')
add_list('Token由独立的用户认证服务（8.159.143.133:8000）签发，使用HS256算法签名')
add_list('前端将Token存储于localStorage（键名：authToken）')
add_list('所有API请求通过Axios拦截器自动携带Authorization: Bearer <token>请求头')
add_list('后端通过PyJWT解码Token，从payload中提取user_id（支持sub、id、email字段）')
add_list('Token无效或过期时返回401 Unauthorized，前端自动跳转登录页')

add_heading('8.2 权限控制', level=2)
add_para('前端通过全局路由中间件（middleware/auth.global.js）实现基于权限的访问控制：')
add_list('权限字段存储于Token payload中：access_system_a ~ access_system_z 和 is_superuser')
add_list('每个功能模块映射到特定的权限字段，如污渍检测模块需要access_system_f权限')
add_list('未授权用户访问受限页面时自动跳转至首页')
add_list('超级管理员（is_superuser=true）可访问所有模块')

doc.add_page_break()

# ============ 9. 部署设计 ============
add_heading('9. 部署设计', level=1)

add_heading('9.1 后端部署', level=2)
add_para('后端采用Docker容器化部署：')
add_list('基础镜像：python:3.12-slim')
add_list('系统依赖：安装libgl1-mesa-glx等OpenCV/YOLO所需的图形库')
add_list('Python依赖：通过requirements.txt安装（使用阿里云PyPI镜像加速）')
add_list('暴露端口：8081')
add_list('启动命令：uvicorn app.main:app --host 0.0.0.0 --port 8081')
add_para('Docker Compose配置：')
add_list('服务名：backend，容器名：stain-backend')
add_list('端口映射：8081:8081')
add_list('卷挂载：./models（YOLO模型文件）、.env（环境变量）、./app（应用代码）')
add_list('重启策略：unless-stopped')

add_heading('9.2 前端部署', level=2)
add_para('前端采用Nuxt.js 3内置服务器部署：')
add_list('开发模式：pnpm dev（默认端口3000）')
add_list('生产构建：nuxt build + node .output/server/index.mjs')
add_list('静态生成：nuxt generate（适用于纯静态部署）')
add_list('CSR模式：ssr: false，所有渲染在客户端完成')
add_list('反向代理：生产环境建议通过Nginx反向代理，配置300秒超时（大图像处理场景）')

add_heading('9.3 环境变量配置', level=2)
add_para('后端环境变量（通过.env文件或Docker环境变量注入）：')
add_table(
    ['变量名', '默认值', '说明'],
    [
        ['APP_NAME', 'Stone Curtain Wall Stain Detection API', '应用名称'],
        ['APP_ENV', 'development', '运行环境'],
        ['CORS_ORIGINS', 'http://localhost:80', '允许的跨域来源'],
        ['OSS_BASE_URL', 'http://8.159.143.133:9000', 'OSS服务地址'],
        ['OSS_UPLOAD_ENDPOINT', '/oss/upload/output', 'OSS上传端点'],
        ['OSS_USERNAME', 'stain-detection', 'OSS用户名'],
        ['OSS_PASSWORD', 'tongji-icw-3567', 'OSS密码'],
        ['MYSQL_HOST', '8.159.143.133', 'MySQL主机'],
        ['MYSQL_PORT', '3306', 'MySQL端口'],
        ['MYSQL_USER', 'stain', 'MySQL用户名'],
        ['MYSQL_PASSWORD', 'stain', 'MySQL密码'],
        ['MYSQL_DATABASE', 'stain', 'MySQL数据库名'],
        ['JWT_SECRET', '', 'JWT签名密钥'],
        ['SYNC_SIZE_THRESHOLD_BYTES', '2097152', '同步处理阈值（2MB）'],
        ['YOLO_MODEL_PATH', 'models/best.pt', 'YOLO模型文件路径'],
        ['YOLO_CONFIDENCE_THRESHOLD', '0.25', 'YOLO置信度阈值'],
        ['CLOUD_MODEL_URL', '', '云端推理API地址'],
        ['CLOUD_MODEL_API_KEY', '', '云端推理API密钥'],
        ['CLOUD_MODEL_IOU_THRESHOLD', '0.7', '云端模型IoU阈值'],
        ['CLOUD_MODEL_IMGSZ', '640', '云端模型输入图像尺寸'],
        ['CLOUD_MODEL_TIMEOUT_SECONDS', '60', '云端推理超时时间'],
    ]
)

doc.add_page_break()

# ============ 10. 附录 ============
add_heading('10. 附录', level=1)

add_heading('10.1 缩略语与术语表', level=2)
add_table(
    ['缩略语/术语', '全称', '说明'],
    [
        ['YOLO', 'You Only Look Once', '实时目标检测深度学习模型系列'],
        ['FastAPI', '-', 'Python异步Web框架'],
        ['Nuxt.js', '-', '基于Vue.js的全栈框架'],
        ['OSS', 'Object Storage Service', '对象存储服务'],
        ['MinIO', '-', '开源对象存储服务（兼容S3协议）'],
        ['JWT', 'JSON Web Token', 'JSON网络令牌，用于身份认证'],
        ['CSR', 'Client-Side Rendering', '客户端渲染模式'],
        ['SSR', 'Server-Side Rendering', '服务端渲染模式'],
        ['API', 'Application Programming Interface', '应用程序编程接口'],
        ['RESTful', 'Representational State Transfer', '表述性状态转移架构风格'],
        ['Pydantic', '-', 'Python数据验证库'],
        ['PyJWT', '-', 'Python JWT编解码库'],
        ['PyMySQL', '-', 'Python MySQL数据库驱动'],
        ['Pillow (PIL)', 'Python Imaging Library', 'Python图像处理库'],
        ['IoU', 'Intersection over Union', '交并比，目标检测评估指标'],
        ['BBox', 'Bounding Box', '边界框，目标检测中的检测框'],
    ]
)

add_heading('10.2 参考文档', level=2)

add_heading('10.2.1 项目文档', level=3)
add_list('POS-石材幕墙污渍检测系统.docx（项目定义说明）')
add_list('作业4_项目章程_石材幕墙污渍检测系统_组长2353740于广淳.docx（项目章程）')
add_list('7_SDS_Intelligent_Curtain_Wall_Stain_Detection.docx（参考SDS文档）')
add_list('SRS_石材幕墙污渍检测系统.pdf（软件需求规格说明）')

add_heading('10.2.2 技术参考', level=3)
add_list('YOLOv8 官方文档：https://docs.ultralytics.com/')
add_list('FastAPI 官方文档：https://fastapi.tiangolo.com/')
add_list('Nuxt.js 3 官方文档：https://nuxt.com/')
add_list('Vue 3 官方文档：https://vuejs.org/')
add_list('Element Plus 官方文档：https://element-plus.org/')
add_list('Docker 官方文档：https://docs.docker.com/')
add_list('MySQL 8.0 参考手册：https://dev.mysql.com/doc/')
add_list('IEEE 1016-2009 软件设计规格说明标准')

# ============ 保存 ============
output_path = r'd:\TongJi\大三下\软件工程管理与经济\期末\StoneCurtainWallStainDetectionSystem\SDS_石材幕墙污渍检测系统_改进版.docx'
doc.save(output_path)
print(f'文档已保存至: {output_path}')
